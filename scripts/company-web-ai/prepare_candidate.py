"""Prepare reviewable server copies; never writes the live server or source snapshot."""
import argparse
import hashlib
import json
import shutil
from pathlib import Path


def replace(text, before, after, count=1):
    if text.count(before)!=count:
        raise RuntimeError('Source drift: expected '+repr(before[:100]))
    return text.replace(before,after)


def prepare(source, destination):
    source,destination=Path(source).resolve(),Path(destination).resolve()
    if destination.exists() or source==destination or source in destination.parents:
        raise RuntimeError('Use a new sibling destination directory')
    shutil.copytree(source,destination)
    manifest=[]
    for server in destination.rglob('20260923-weekly-server.py'):
        folder=server.parent
        shutil.copyfile(Path(__file__).with_name('hub_ai_queue.py'),folder/'hub_ai_queue.py')
        shutil.copyfile(Path(__file__).with_name('hub_poster.py'),folder/'hub_poster.py')
        # Generated once on the build machine; server needs only Pillow and ZIP/XML.
        from pptx import Presentation
        from pptx.util import Inches
        presentation=Presentation();presentation.slide_width=Inches(13.333333);presentation.slide_height=Inches(7.5)
        presentation.slides.add_slide(presentation.slide_layouts[6]);presentation.save(folder/'hub-poster-base.pptx')
        text=server.read_text(encoding='utf-8')
        start=text.index('def polish_deepseek(');end=text.index('\n\nclass Job(',start)
        text=text[:start]+'''def polish_deepseek(key, text, system=POLISH_SYSTEM):
    from hub_ai_queue import provider, HubAiError
    try:
        result = provider(key, text, system)
    except HubAiError as error:
        raise HTTPException(503, str(error)) from None
    if system == POLISH_SYSTEM:
        validate_polish_format(result['text'])
        result['text'] = restrain_polish_emphasis(result['text'])
    return result
'''+text[end:]
        text=replace(text,"    app.state.polisher = polish_deepseek", "    import sys\n    sys.path.insert(0, str(Path(__file__).parent))\n    from hub_ai_queue import attach_fastapi\n    app.state.hub_ai_queue = attach_fastapi(app, Path(settings['data_dir']) / 'hub-ai')\n    app.state.polisher = polish_deepseek")
        text=replace(text,"        key = settings.get('deepseek_api_key', '')\n        if not key:\n            raise HTTPException(503, 'DeepSeek 尚未配置，请检查服务配置')","        key = 'CodeAgent'")
        text=replace(text,"        templates.validate(body.style, body.scope)","        templates.validate(body.style, body.scope)\n        if body.style not in ('A1','B1','C1'): raise HTTPException(422, '请选择内置海报模板')")
        # Keep legacy rows unchanged, but allow users to clone their source into templates.
        text=replace(text,"AND state IN ('queued','submitted','running','needs_attention','parked') ORDER BY created DESC LIMIT 1", "AND phase LIKE 'codeagent%' AND state IN ('queued','submitted','running','needs_attention','parked') ORDER BY created DESC LIMIT 1",count=2)
        text=replace(text,"if old['state']!='failed':raise HTTPException", "if old['state']!='failed' and (old['phase'].startswith('codeagent') or old['state']=='complete'):raise HTTPException")
        text=replace(text,"old['materials'],old['style'],now,now", "old['materials'],old['style'] if old['style'] in ('A1','B1','C1') else 'C1',now,now")
        text=replace(text,'    def create_job(body: Job, request: Request):\n        me = user(request)',"    def create_job(body: Job, request: Request):\n        me = user(request)\n        if not re.fullmatch('[a-f0-9]{64}',request.headers.get('x-hub-ai-client','')): raise HTTPException(503,'请在公司 Hub 的团队页面生成海报')")
        text=replace(text,'    def retry_job(jid: str, body: RetryJob, request: Request):\n        user(request)',"    def retry_job(jid: str, body: RetryJob, request: Request):\n        user(request)\n        if not re.fullmatch('[a-f0-9]{64}',request.headers.get('x-hub-ai-client','')): raise HTTPException(503,'请在公司 Hub 的团队页面重试')")
        text=replace(text,"references=[templates.reference(body.style)] if body.style in templates.BY_ID else [])","references=[])")
        text=replace(text,"VALUES(?,?,?,?,?,?,'queued',?,?,?,?,?,?,?,?)","VALUES(?,?,?,?,?,?,'running',?,?,?,?,?,?,?,?)",count=2)
        text=replace(text,'            return public_job(db, job_record(db, jid))',"            db.execute(\"UPDATE jobs SET phase='codeagent_content' WHERE id=?\",(jid,))\n            saved=public_job(db, job_record(db, jid))\n        app.state.hub_posters.submit(jid)\n        return saved")
        text=replace(text,'            return public_job(db,job_record(db,new_id))',"            db.execute(\"UPDATE jobs SET phase='codeagent_content' WHERE id=?\",(new_id,))\n            saved=public_job(db,job_record(db,new_id))\n        app.state.hub_posters.submit(new_id)\n        return saved")
        text=replace(text,'    return app',"    from hub_poster import Posters\n    app.state.hub_posters = Posters(store, app.state.hub_ai_queue)\n    return app")
        text=replace(text,"    def work(request: Request, template_refs: bool = False):\n        bridge(request)","    def work(request: Request, template_refs: bool = False):\n        bridge(request)\n        return {'jobs': []}  # Retired external-image connector; existing records remain intact.")
        text=replace(text,"        message=messages.get(row['error_code'], '生成遇到问题，请检查生成服务连接' if row['error_code'] else '')","        message=(row['error_code'] if row['phase'].startswith('codeagent') else messages.get(row['error_code'], '生成遇到问题，请检查生成服务连接' if row['error_code'] else ''))")
        server.write_text(text,encoding='utf-8')
        summary=folder/'20260924-weekly-summary.py';text=summary.read_text(encoding='utf-8')
        text=replace(text,"        key=app.state.settings.get('deepseek_api_key','')\n        if not key:raise HTTPException(503,'DeepSeek 尚未配置，请检查服务配置')","        key='CodeAgent'")
        summary.write_text(text.replace('DeepSeek 未返回完整摘要','CodeAgent 未返回完整摘要'),encoding='utf-8')
        web=folder/'20260923-weekly-web.js';text=web.read_text(encoding='utf-8')
        text=text.replace(',45000)',',660000)')
        web.write_text(text.replace('DeepSeek','CodeAgent').replace('图片完成后会自动转为可编辑 PPT。','CodeAgent 整理后按模板生成图片与可编辑 PPT。'),encoding='utf-8')
        html=folder/'20260923-weekly-web.html';text=html.read_text(encoding='utf-8')
        text=text.replace('DeepSeek','CodeAgent').replace('图片完成后自动转字，文字可编辑；复杂图形保留为背景。','模板直接生成可编辑文字与版式，无需识图转字。').replace('生成需要几分钟。图片返回后自动核对文字与素材，疑点仅作提醒，可在 PPT 中修改。离开页面后任务仍会继续。','CodeAgent 通常需要1至3分钟。生成后可点击事实核对，确认内容后再发布；请保持 Hub 运行。')
        html.write_text(text,encoding='utf-8')
        collaboration=folder/'20260928-weekly-collaboration.js'
        collaboration.write_text(replace(collaboration.read_text(encoding='utf-8'),"const collaborationVersion='2026.09.28.2'","const collaborationVersion='4.1.1 · CodeAgent 模板版'"),encoding='utf-8')
        progress=folder/'20260925-poster-progress.js';text=progress.read_text(encoding='utf-8')
        progress.write_text(text.replace('DeepSeek 正在对照素材核对','CodeAgent 正在对照素材核对').replace('已用原素材和原模板重新排队，失败记录保留。','已用原素材创建模板任务，原记录保留。')+'\n'+Path(__file__).with_name('hub_template_ui.js').read_text(encoding='utf-8'),encoding='utf-8')
        # Only the three fixed templates are exposed. Old records keep their original styles.
        template_file=folder/'20260924-weekly-templates.py'
        text=template_file.read_text(encoding='utf-8')
        text=replace(text,"    if code in BY_ID and scope!='group':raise HTTPException(422,'这12套模板用于小组周报，个人周报请选择自由创作')","    if code not in ('A1','B1','C1'):raise HTTPException(422,'请选择内置海报模板')")
        template_file.write_text(text,encoding='utf-8')
        import sys,importlib.util
        sys.path.insert(0,str(folder))
        spec=importlib.util.spec_from_file_location('hub_poster_build',folder/'hub_poster.py');poster=importlib.util.module_from_spec(spec);spec.loader.exec_module(poster)
        catalog=[];assets=folder/'20260924-template-library'
        sample={'sections':[{'heading':'重点方向 '+str(i+1),'progress':'填写本周已验证的进展、关键条件和当前状态。','next':'填写下一步计划与待确认事项。'} for i in range(4)],'members':'成员甲、成员乙、成员丙'}
        for code,title in poster.TEMPLATES.items():
            png=assets/('hub-template-'+code+'.png');ppt=assets/('hub-template-'+code+'.pptx')
            poster.render(sample,code,'团队周报模板','YYYY-MM-DD',png,ppt)
            catalog.append(dict(code=code,title=title,group=code[0],density='1至4项重点 · 固定版式',description=title+'，CodeAgent只提炼内容，模板确定性排版。',filename=png.name,sha256=hashlib.sha256(png.read_bytes()).hexdigest()))
        (assets/'20260924-template-catalog.json').write_text(json.dumps(catalog,ensure_ascii=False),encoding='utf-8')
        audit=folder/'20260925-weekly-audit.py';text=audit.read_text(encoding='utf-8')
        exclusion="job NOT IN (SELECT id FROM jobs WHERE phase LIKE 'codeagent%') AND "
        text=replace(text,"WHERE state IN ('reading','checking') AND lease_until<? AND attempts>=3", "WHERE "+exclusion+"state IN ('reading','checking') AND lease_until<? AND attempts>=3")
        text=replace(text,"WHERE state='queued' OR (state IN ('reading','checking') AND lease_until<? AND attempts<3) ORDER BY created LIMIT 1", "WHERE "+exclusion+"(state='queued' OR (state IN ('reading','checking') AND lease_until<? AND attempts<3)) ORDER BY created LIMIT 1")
        text=replace(text,'    def request_audit(jid:str,request:Request):\n        user(request)','    def request_audit(jid:str,request:Request):\n        user(request)\n        app.state.hub_ai_queue.bind_reference(jid)\n        if app.state.hub_posters.audit(jid):\n            with store.db() as db: return public(db,jid)')
        text=replace(text,"            key=app.state.settings.get('deepseek_api_key','')\n            if not key:raise HTTPException(503,'DeepSeek 尚未配置，核对未完成')","            key='CodeAgent'")
        # Both reviewer calls are within the same persisted owner binding, including OCR callbacks.
        a=text.index('            raw=app.state.poster_auditor');b=text.index("            message='发现",a)
        block=text[a:b];text=text[:a]+'            with app.state.hub_ai_queue.reference_scope(jid):\n'+''.join('    '+line if line.strip() else line for line in block.splitlines(True))+text[b:]
        audit.write_text(text,encoding='utf-8')
    style=destination/'style/stylemap_server.py';text=style.read_text(encoding='utf-8')
    shutil.copyfile(Path(__file__).with_name('hub_ai_queue.py'),style.parent/'hub_ai_queue.py')
    text=replace(text,'        self.pool = ThreadPoolExecutor(max_workers=16)','        self.pool = ThreadPoolExecutor(max_workers=16)\n        from hub_ai_queue import Queue\n        self.hub_queue = Queue(os.path.join(cache_dir, "hub-ai"))')
    text=replace(text,'        self.pool.submit(self._run, h, system, user, check)','        from contextvars import copy_context\n        self.pool.submit(copy_context().run, self._run, h, system, user, check)')
    text=replace(text,'        for attempt in range(2):','        for attempt in range(1):')
    a=text.index('    def call(self, system, user):');b=text.index('\n\ndef _need',a)
    text=text[:a]+'''    def call(self, system, user):
        from hub_ai_queue import provider
        txt = provider('', user, system)['text'].strip()
        if txt.startswith('```'):
            txt = re.sub(r'^```(?:json)?\\s*|\\s*```$', '', txt)
        return json.loads(txt)
'''+text[b:]
    text=replace(text,'    return H','    from hub_ai_queue import wrap_http_handler\n    return wrap_http_handler(H, ai.hub_queue, BASE)')
    text=replace(text,'AI(cfg["deepseek_key"], cfg.get("model", "deepseek-v4-pro"),','AI("", "CodeAgent",')
    style.write_text(text,encoding='utf-8')
    for file in destination.rglob('*'):
        if not file.is_file() or file.suffix not in ('.py','.js','.json','.png','.pptx','.html'):continue
        original=source/file.relative_to(destination)
        if not original.exists() or original.read_bytes()!=file.read_bytes():
            if file.suffix=='.py':compile(file.read_text(encoding='utf-8'),str(file),'exec')
            relative=file.relative_to(destination).as_posix()
            roots={'weekly/':'C:/TeamWeekly/','ai-team/weekly_engine/':'C:/AITeamWeekly/app/weekly_engine/','reports/template/weekly_engine/':'C:/WeeklyPlatform/app/template/weekly_engine/','style/':'C:/stylemap/app/'}
            remote=next((target+relative[len(prefix):] for prefix,target in roots.items() if relative.startswith(prefix)),None)
            if remote is None:raise RuntimeError('No reviewed remote mapping for '+relative)
            manifest.append({'path':relative,'remote':remote,'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'before_sha256':hashlib.sha256(original.read_bytes()).hexdigest() if original.exists() else None})
    (destination/'20261024-codeagent-candidate.json').write_text(json.dumps({'changes':manifest,'deployed':False,'posterGeneration':'CodeAgent content + three deterministic PNG/PPT templates; external connector retired'},ensure_ascii=False,indent=2),encoding='utf-8')
    return manifest


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('source');parser.add_argument('destination');args=parser.parse_args()
    print(json.dumps({'changed':len(prepare(args.source,args.destination)),'deployed':False}))
