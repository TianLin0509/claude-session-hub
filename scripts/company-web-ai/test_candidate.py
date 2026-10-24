"""Exercise each real copied weekly app with synthetic records and a fake worker.
Network model APIs are disabled; this tests routing/rendering, not model quality.
"""
import argparse
import concurrent.futures
import importlib.util
import io
import json
import sys
import tempfile
import time
import urllib.request
import uuid
from pathlib import Path
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation


def test(folder):
    folder=Path(folder).resolve();sys.path.insert(0,str(folder))
    spec=importlib.util.spec_from_file_location('weekly_test',folder/'20260923-weekly-server.py');module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module)
    external=[]
    def no_network(*args,**kwargs):
        external.append(True);raise AssertionError('External model API attempted')
    urllib.request.urlopen=no_network
    with tempfile.TemporaryDirectory() as root:
        app=module.create_app({'public_origin':'http://testserver','connector_token':'x'*40,'data_dir':root,'timezone':'Asia/Shanghai','deepseek_api_key':'old-key-must-not-be-used'})
        store=app.state.store;store.add_member('测试甲');store.add_member('测试乙');store.ensure_weeks()
        week=store.current_week()
        with store.db() as db:member=store.members(db)[0]['id']
        headers={'Origin':'http://testserver','X-Weekly-Request':'1','X-Hub-AI-Client':'a'*64}
        client=TestClient(app);worker=TestClient(app)
        def claim():
            for _ in range(150):
                response=worker.post('/api/hub-ai/claim',headers=headers,json={});assert response.status_code==200,response.text
                job=response.json()['job']
                if job:return job
                time.sleep(.02)
            raise AssertionError('No CodeAgent task')
        def complete(job,text):
            response=worker.post('/api/hub-ai/complete',headers=headers,json={'id':job['id'],'receipt':job['receipt'],'text':text});assert response.status_code==200,response.text
        with concurrent.futures.ThreadPoolExecutor() as pool:
            future=pool.submit(client.post,'/api/summary',headers=headers,json={'week':week,'materials':'测试甲完成基线测试；测试乙计划复核。','request_id':str(uuid.uuid4())})
            job=claim();assert '测试甲' in job['text'];assert worker.post('/api/hub-ai/claim',headers={**headers,'X-Hub-AI-Client':'b'*64},json={}).json()['job'] is None
            complete(job,'测试甲完成基线测试，测试乙计划复核。')
            response=future.result(5);assert response.status_code==200,response.text;assert response.json()['model']=='CodeAgent'
        catalog=client.get('/api/templates').json();assert len(catalog)==3
        for template in catalog:
            sample=client.get(template['image_url']);assert sample.status_code==200;assert Image.open(io.BytesIO(sample.content)).size==(1600,900)
        for code in ('A1','B1','C1'):
            request={'week':week,'materials':'测试甲完成基线测试，条件为单小区；测试乙计划复核。','style':code,'request_id':str(uuid.uuid4()),'auto_ppt':True}
            response=client.post('/api/jobs',headers=headers,json=request);assert response.status_code==200,response.text;jid=response.json()['id']
            task=claim();assert '只负责提炼内容' in task['system']
            assert client.get('/api/bridge/work',headers={'Authorization':'Bearer '+'x'*40}).json()=={'jobs':[]}
            content={'sections':[{'heading':'基线测试','progress':'测试甲完成单小区基线测试。','next':'测试乙计划复核。'}],'members':'测试甲、测试乙'}
            complete(task,json.dumps(content,ensure_ascii=False))
            for _ in range(150):
                jobs=client.get('/api/jobs',params={'week':week},headers=headers).json();done=next(x for x in jobs if x['id']==jid)
                if done['state'] in ('complete','failed'):break
                time.sleep(.03)
            assert done['state']=='complete',done
            image=client.get(done['image_url'],headers=headers);assert image.status_code==200;assert Image.open(io.BytesIO(image.content)).size==(1600,900)
            ppt=client.get(done['ppt']['url'],headers=headers);assert ppt.status_code==200
            presentation=Presentation(io.BytesIO(ppt.content));assert len(presentation.slides)==1
            text='\n'.join(shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame)
            assert '单小区' in text and '测试乙计划复核' in text
            assert client.post('/api/jobs',headers=headers,json=request).json()['id']==jid
            if code=='C1':
                audit=client.post('/api/jobs/'+jid+'/audit',headers=headers,json={});assert audit.status_code==200,audit.text
                old_claim=client.post('/api/bridge/audit/claim',headers={'Authorization':'Bearer '+'x'*40},json={'worker':str(uuid.uuid4())})
                assert old_claim.status_code==200 and old_claim.json()['job'] is None,old_claim.text
                task=claim();assert '事实核对' in task['system'];complete(task,'{"issues":[]}')
                for _ in range(100):
                    with store.db() as db:state=db.execute('SELECT state FROM poster_audits WHERE job=?',(jid,)).fetchone()[0]
                    if state=='complete':break
                    time.sleep(.02)
                assert state=='complete'
                # Simulate a dead process's pending audit, then reconstruct the service.
                with store.db(True) as db:db.execute("UPDATE poster_audits SET state='checking' WHERE job=?",(jid,))
                from hub_poster import Posters
                app.state.hub_posters=Posters(store,app.state.hub_ai_queue)
                with store.db() as db:assert db.execute('SELECT state FROM poster_audits WHERE job=?',(jid,)).fetchone()[0]=='failed'
                assert client.get(done['ppt']['url'],headers=headers).status_code==200
                assert client.post('/api/jobs/'+jid+'/audit',headers=headers,json={}).status_code==200
                complete(claim(),'{"issues":[]}')
                for _ in range(100):
                    with store.db() as db:state=db.execute('SELECT state FROM poster_audits WHERE job=?',(jid,)).fetchone()[0]
                    if state=='complete':break
                    time.sleep(.02)
                assert state=='complete'
        # Legacy active records remain untouched and can be cloned into a fixed template.
        with store.db(True) as db:
            db.execute("UPDATE jobs SET phase='waiting_image',state='parked',style='F1' WHERE id=?",(jid,))
        retry={'request_id':str(uuid.uuid4())}
        cloned=client.post('/api/jobs/'+jid+'/retry',headers=headers,json=retry)
        assert cloned.status_code==200,cloned.text
        new_id=cloned.json()['id'];assert new_id!=jid
        assert client.post('/api/jobs/'+jid+'/retry',headers=headers,json=retry).json()['id']==new_id
        with store.db() as db:
            assert db.execute('SELECT state FROM jobs WHERE id=?',(jid,)).fetchone()[0]=='parked'
            assert db.execute('SELECT style FROM jobs WHERE id=?',(new_id,)).fetchone()[0]=='C1'
        complete(claim(),json.dumps(content,ensure_ascii=False))
        for _ in range(100):
            with store.db() as db:state=db.execute('SELECT state FROM jobs WHERE id=?',(new_id,)).fetchone()[0]
            if state=='complete':break
            time.sleep(.03)
        assert state=='complete'
        assert external==[]
        print(json.dumps({'app':str(folder),'passed':True,'templates':3,'externalAiCalls':0,'checks':['summary queue + owner isolation','real PNG + editable PPT','request idempotency','external connector retired','direct template fact check without OCR']},ensure_ascii=False))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('folder');args=p.parse_args();test(args.folder)
